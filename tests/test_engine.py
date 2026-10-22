"""Прогон всех действий движка на сгенерированных файлах: python tests/test_engine.py"""
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import engine  # noqa: E402


def make_samples(d: Path) -> dict:
    from PIL import Image, ImageDraw
    import qrcode
    import pymupdf

    s = {}
    im = Image.new("RGBA", (640, 400), (255, 140, 0, 255))
    ImageDraw.Draw(im).ellipse((100, 50, 400, 350), fill=(30, 160, 60, 128))
    s["png"] = d / "фото.png"; im.save(s["png"])
    exif = Image.Exif(); exif[0x010F] = "TestCam"
    s["jpg"] = d / "снимок.jpg"; im.convert("RGB").save(s["jpg"], exif=exif)
    s["qr"] = d / "qr.png"; qrcode.make("https://example.com/mandarin").save(s["qr"])
    frames = [Image.new("RGB", (120, 80), c) for c in ("red", "green", "blue")]
    s["gif"] = d / "anim.gif"; frames[0].save(s["gif"], save_all=True, append_images=frames[1:], duration=200)

    doc = pymupdf.open()
    for i in range(3):
        page = doc.new_page()
        page.insert_text((72, 72), f"Mandarin test page {i + 1}")
        page.insert_image(pymupdf.Rect(72, 100, 300, 300), filename=str(s["jpg"]))
    s["pdf"] = d / "док.pdf"; doc.save(s["pdf"]); doc.close()

    ff = engine.ffmpeg_exe()
    import docx
    dd = docx.Document(); dd.add_heading("Отчёт", 1); dd.add_paragraph("Проверка Mandarin.")
    s["docx"] = d / "отчёт.docx"; dd.save(s["docx"])

    s["wav"] = d / "звук.wav"
    subprocess.run([ff, "-loglevel", "error", "-f", "lavfi", "-i", "sine=frequency=440:duration=2",
                    str(s["wav"])], check=True)
    s["mp4"] = d / "видео.mp4"
    subprocess.run([ff, "-loglevel", "error", "-f", "lavfi", "-i", "testsrc=duration=2:size=320x240:rate=15",
                    "-f", "lavfi", "-i", "sine=duration=2", "-c:v", "libx264", "-pix_fmt", "yuv420p",
                    "-shortest", str(s["mp4"])], check=True)
    return s


def main():
    d = Path(tempfile.mkdtemp(prefix="mandarin-test-"))
    s = make_samples(d)
    cases = [
        ([s["png"]], "convert"), ([s["png"]], "tools"),
        ([s["jpg"]], "convert"), ([s["jpg"]], "tools"),
        ([s["gif"]], "convert"),
        ([s["pdf"]], "convert"), ([s["pdf"]], "tools"),
        ([s["pdf"], s["png"]], "convert"),
        ([s["wav"]], "convert"), ([s["wav"]], "tools"),
        ([s["mp4"]], "convert"), ([s["mp4"]], "tools"),
        ([s["docx"]], "convert"),
    ]
    ok = fail = skip = 0
    for paths, mode in cases:
        cat = engine.common_category(paths)
        for a in engine.actions_for(paths, mode):
            if a.ui or a.ask:
                continue  # окна и действия с параметрами — в test_organizer / test_pdftools / test_viewer
            if a.id == "img_qr" and paths[0] != s["qr"]:
                paths_run = [s["qr"]]
            else:
                paths_run = paths
            try:
                r = a.run(paths_run)
                bad = [o for o in r.outputs if not o.exists() or (o.is_file() and o.stat().st_size == 0)]
                assert not bad, f"пустые/нет файлов: {bad}"
                assert r.outputs or r.clipboard, "нет результата"
                ok += 1
                info = r.clipboard[:40].replace("\n", " ") if r.clipboard else ", ".join(o.name for o in r.outputs[:2])
                print(f"  ok   {cat:6} {mode:7} {a.label:9} → {info} {r.message}")
            except engine.EngineError as e:
                expected = ("LibreOffice", "компактный", "текстовый слой", "нет цифровых подписей")
                if any(x in str(e) for x in expected):
                    skip += 1
                    print(f"  skip {cat:6} {mode:7} {a.label:9} ({e})")
                else:
                    fail += 1
                    print(f"  FAIL {cat:6} {mode:7} {a.label:9} {e}")
            except Exception as e:
                fail += 1
                print(f"  FAIL {cat:6} {mode:7} {a.label:9} {type(e).__name__}: {e}")

    # без офисных пакетов: документы скрыты, PDF→DOCX остаётся (pdf2docx)
    real = engine.soffice_exe
    engine.soffice_exe = lambda: None
    try:
        assert engine.actions_for([s["docx"]], "convert") == [], "документы без офиса должны скрываться"
        ids = [a.id for a in engine.actions_for([s["pdf"]], "convert")]
        assert "pdf_docx" in ids, ids
        r = engine.pdf_to_docx([s["pdf"]])
        assert r.outputs[0].stat().st_size > 0
        print("  ok   без LibreOffice: документы скрыты, PDF→DOCX работает")
    finally:
        engine.soffice_exe = real

    # метаданные действительно удалены
    from PIL import Image
    stripped = next(d.glob("снимок (без метаданных)*.jpg"))
    assert not Image.open(stripped).getexif(), "EXIF не удалён"
    assert Image.open(s["jpg"]).getexif(), "оригинал изменён!"
    print(f"\nИтого: {ok} ok, {fail} fail, {skip} skip (нет LibreOffice / уже компактный)")
    shutil.rmtree(d)
    sys.exit(1 if fail else 0)


if __name__ == "__main__":
    main()
